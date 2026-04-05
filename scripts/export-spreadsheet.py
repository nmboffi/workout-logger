#!/usr/bin/env python3
"""Export workout logger data back to SBS-format Excel spreadsheet.

Usage: python export-spreadsheet.py <data_json> [output.xlsx]
  data_json: exported JSON from the app (via Settings > Export Data)
  output.xlsx: output file (default: SBS_Hypertrophy_export.xlsx)

If data_json is "program", reads from data/program.json and generates the
spreadsheet from the imported program data (no workout logs).
"""

import json
import sys
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Install openpyxl: pip install openpyxl")
    sys.exit(1)

# Load program data
with open("data/program.json") as f:
    program = json.load(f)

def parse_app_data(path):
    """Parse exported app data JSON, or return None if using program.json only."""
    if path == "program":
        return None
    with open(path) as f:
        return json.load(f)

def create_workbook(app_data=None):
    wb = openpyxl.Workbook()
    config = program["config"]

    # Colors
    header_fill = PatternFill("solid", fgColor="2A2118")
    card_fill = PatternFill("solid", fgColor="362D24")
    amber_font = Font(color="F5A623", bold=True)
    white_font = Font(color="FFF8F0", bold=True)
    muted_font = Font(color="8C7D6D")
    body_font = Font(color="C4B5A5")
    thin_border = Border(
        bottom=Side(style="thin", color="4A3F33"),
    )

    # Get TMs - from app data or program defaults
    tms = {}
    if app_data and "training_maxes" in app_data:
        tms = app_data["training_maxes"]
    else:
        for lift in config["mainLifts"]:
            tms[lift["name"]] = lift["trainingMax"]
        for lift in config["auxiliaries"]:
            tms[lift["name"]] = lift["trainingMax"]

    schedule_type = "4x"
    if app_data and "schedule_type" in app_data:
        schedule_type = app_data["schedule_type"]

    # Get workout logs
    logs = []
    if app_data and "workout_logs" in app_data:
        logs = app_data["workout_logs"]
    logs_by_week_day = {}
    for log in logs:
        key = (log.get("weekNumber", 0), log.get("dayIndex", 0))
        logs_by_week_day[key] = log

    # ========== Quick Setup sheet ==========
    ws = wb.active
    ws.title = "Quick Setup"
    ws.sheet_properties.tabColor = "F5A623"

    ws["A1"] = "Rounding"
    ws["A2"] = config["rounding"]

    # Main lifts
    ws["C4"] = "Main Lifts"
    ws["C4"].font = white_font
    ws["D4"] = "Maxes"
    ws["D4"].font = muted_font
    ws["E4"] = "Single @8 %"
    ws["E4"].font = muted_font

    for i, lift in enumerate(config["mainLifts"]):
        row = 5 + i
        ws.cell(row=row, column=3, value=lift["name"]).font = body_font
        ws.cell(row=row, column=4, value=tms.get(lift["name"], lift["trainingMax"])).font = amber_font
        ws.cell(row=row, column=5, value=lift["singleAt8Pct"]).font = body_font

    # Auxiliaries
    aux_start = 5 + len(config["mainLifts"]) + 2
    ws.cell(row=aux_start - 1, column=3, value="Variations").font = white_font
    ws.cell(row=aux_start - 1, column=4, value="Maxes").font = muted_font

    for i, aux in enumerate(config["auxiliaries"]):
        row = aux_start + i
        ws.cell(row=row, column=2, value=aux["slot"]).font = muted_font
        ws.cell(row=row, column=3, value=aux["name"]).font = body_font
        ws.cell(row=row, column=4, value=tms.get(aux["name"], aux["trainingMax"])).font = amber_font
        ws.cell(row=row, column=5, value=aux["singleAt8Pct"]).font = body_font

    # Accessory pools
    pool_start = aux_start + len(config["auxiliaries"]) + 2
    pools = config["accessoryPools"]
    col = 2
    for pool_name, exercises in pools.items():
        ws.cell(row=pool_start, column=col, value=pool_name).font = white_font
        for j, ex in enumerate(exercises):
            ws.cell(row=pool_start + 1 + j, column=col, value=ex).font = body_font
        col += 1

    # Autoregulation table
    autoreg_start = pool_start + max(len(v) for v in pools.values()) + 3
    ws.cell(row=autoreg_start, column=7, value="When above or below rep target").font = white_font
    headers = ["Sets", "Below by 2+", "Below by 1", "Hit target", "Beat by 1", "Beat by 2", "Beat by 3", "Beat by 4"]
    for i, h in enumerate(headers):
        ws.cell(row=autoreg_start + 1, column=7 + i if i == 0 else 8 + i - 1, value=h).font = muted_font

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 16

    # ========== Program sheets (4x, (3+1)x) ==========
    for sheet_name in [schedule_type]:
        template = program["templates"].get(sheet_name, [])
        if not template:
            continue

        ws = wb.create_sheet(title=sheet_name)
        ws.sheet_properties.tabColor = "F5A623"

        week_schedule = program["weekSchedule"]
        num_weeks = len(week_schedule)

        # Column layout: col A = exercise names, then per week: Weight, Reps, RepOut, Sets, LastSet, Video, Notes
        cols_per_week = 7

        # Header row
        for wi, week in enumerate(week_schedule):
            start_col = 2 + wi * cols_per_week
            ws.cell(row=1, column=start_col, value=f"Week {week['weekNumber']}").font = white_font
            col_headers = ["Weight", "Reps/Set", "Rep Out", "Sets", "Last Set", "Video", "Notes"]
            for ci, h in enumerate(col_headers):
                ws.cell(row=2, column=start_col + ci, value=h).font = muted_font

        # Fill exercises per day
        row = 3
        for day in template:
            ws.cell(row=row, column=1, value=day["label"]).font = white_font
            row += 1

            for ex in day["exercises"]:
                is_main = ex["category"] == "main"

                if is_main:
                    # TM row
                    ws.cell(row=row, column=1, value=f"{ex['name']} TM").font = muted_font
                    for wi, week in enumerate(week_schedule):
                        start_col = 2 + wi * cols_per_week
                        wc = week["exerciseConfigs"].get(ex["name"])
                        if wc:
                            tm = tms.get(ex["name"], 0)
                            ws.cell(row=row, column=start_col, value=round(tm, 2)).font = body_font
                            ws.cell(row=row, column=start_col + 1, value="single @8").font = muted_font
                    row += 1

                    # Working set row
                    ws.cell(row=row, column=1, value=ex["name"]).font = amber_font
                    for wi, week in enumerate(week_schedule):
                        start_col = 2 + wi * cols_per_week
                        wc = week["exerciseConfigs"].get(ex["name"])
                        if wc:
                            tm = tms.get(ex["name"], 0)
                            rounding = config["rounding"]
                            working_wt = round(tm * wc["intensity"] / rounding) * rounding
                            ws.cell(row=row, column=start_col, value=working_wt).font = amber_font
                            ws.cell(row=row, column=start_col + 1, value=wc["reps"]).font = body_font
                            ws.cell(row=row, column=start_col + 2, value=wc["repOutTarget"]).font = body_font
                            ws.cell(row=row, column=start_col + 3, value=wc["sets"]).font = body_font

                        # Fill in logged data if available
                        log = logs_by_week_day.get((week["weekNumber"], template.index(day)))
                        if log:
                            log_ex = next((e for e in log.get("exercises", []) if e.get("exerciseName") == ex["name"]), None)
                            if log_ex and log_ex.get("repsOnLastSet") is not None:
                                ws.cell(row=row, column=start_col + 4, value=log_ex["repsOnLastSet"]).font = Font(color="5EBB7A", bold=True)
                            if log_ex and log_ex.get("notes"):
                                ws.cell(row=row, column=start_col + 6, value=log_ex["notes"]).font = muted_font
                    row += 1
                else:
                    # Accessory row
                    ws.cell(row=row, column=1, value=ex["name"]).font = body_font
                    row += 1

            row += 1  # blank row between days

        # Set column widths
        ws.column_dimensions["A"].width = 28
        for wi in range(num_weeks):
            for ci in range(cols_per_week):
                col_letter = get_column_letter(2 + wi * cols_per_week + ci)
                ws.column_dimensions[col_letter].width = 12

    # ========== Version Info sheet ==========
    ws = wb.create_sheet(title="Export Info")
    ws["A1"] = "Exported from Workout Logger"
    ws["A1"].font = white_font
    ws["A2"] = f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A2"].font = body_font
    ws["A3"] = f"Schedule: {schedule_type}"
    ws["A3"].font = body_font
    ws["A4"] = f"Program: {program['name']}"
    ws["A4"].font = body_font

    if app_data and "current_week" in app_data:
        ws["A5"] = f"Current Week: {app_data['current_week']}"
        ws["A5"].font = body_font

    ws["A7"] = "Training Maxes"
    ws["A7"].font = white_font
    row = 8
    for name, tm in sorted(tms.items()):
        ws.cell(row=row, column=1, value=name).font = body_font
        ws.cell(row=row, column=2, value=round(tm, 1)).font = amber_font
        row += 1

    return wb

def main():
    data_source = sys.argv[1] if len(sys.argv) > 1 else "program"
    output = sys.argv[2] if len(sys.argv) > 2 else "SBS_Hypertrophy_export.xlsx"

    app_data = parse_app_data(data_source)
    wb = create_workbook(app_data)
    wb.save(output)
    print(f"Exported to {output}")

    # Summary
    config = program["config"]
    tms = {}
    if app_data and "training_maxes" in app_data:
        tms = app_data["training_maxes"]
    else:
        for lift in config["mainLifts"]:
            tms[lift["name"]] = lift["trainingMax"]
        for lift in config["auxiliaries"]:
            tms[lift["name"]] = lift["trainingMax"]

    print(f"Training Maxes:")
    for name, tm in sorted(tms.items()):
        print(f"  {name}: {tm:.1f}")

if __name__ == "__main__":
    main()
